#!/usr/bin/env python3
"""Derive the golden key from the shipped inputs alone, and check it against solution/files (GLD-11).

Author tool, not on the reward path. It reads only environment/input/ and applies the disclosed rules:
  R1  (About sheet) the decisions are the lines of edit_thread.txt that name an edit, in message order;
      a reply's quoted lines (">") are the earlier message, not new decisions; an edit's latest decision is
      its call; a call is ACCEPTED (goes in, with the number the decision gives, else the edit's number as it
      last stood), DEFERRED (does not go in, stays open against the passage) or WITHDRAWN (does not go in,
      nothing stays open); a decision returning an edit to an earlier call makes it stand as that call did;
      a decision given as the same call as another makes the same kind of call at this edit's own number
  L0  (layout spec) the ledger lists flags in the order they were raised; `draft_line` (where the flag
      starts in the draft) gives manuscript order; a passage flagged in more than one stretch is one
      passage whose stretches run on from each other, and its page is the page its first line falls on
  L1  30 lines a page; passages run back to back in manuscript order from line 1
  L2  a passage's length is its drafted lines plus the line change of every edit that stands accepted for
      it; a deferred or withdrawn edit changes nothing
  L3  a passage never falls below 1 line (held at 1, floored), and that line still pushes every later
      passage down
  L4  verdict priority: floored > carrying a deferred edit > first and last line on different pages
      (STRADDLES_BREAK) > all on one page (WHOLLY_ON_PAGE)
  L5  final_page_count is the page the manuscript's last line falls on

DECISIONS below is the declared meaning of every wording the Log uses (the interpretation key a reviewer
audits); derive() asserts every Log decision is covered. The same solver, with one switch per wrong
reading, is imported by tests/discrimination.py and tools/triage_trials.py.

It also asserts the fixture's design invariants (README section 5), that the Draft pages sheet the requester
vouches for is the per-flag draft layout, that replies quote earlier decisions, and that every pin in
tests/verifier.json equals the derived gold. Needs openpyxl (the task image has it).

  python3 tests/derive_gold.py            compare with solution/files, exit 1 on any difference
  python3 tests/derive_gold.py --write    rewrite solution/files/page_register.csv and results.json
"""
from __future__ import annotations

import csv
import io
import json
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
INP = ROOT / "environment" / "input"
GOLD = ROOT / "solution" / "files"
LINES_PER_PAGE = 30
HEADER = ["passage_id", "page_number", "verdict"]

#: decision wording -> (kind, lines). kind: accept | defer | withdraw | revert:<reader> | ditto | same:<code>.
#: lines: the signed line change the decision gives, or None when it gives none.
DECISIONS = {
    "Yes, take it.": ("accept", None),
    "Agreed.": ("accept", None),
    "Fine as proposed.": ("accept", None),
    "Go ahead.": ("accept", None),
    "Yes, cut the whole sequence.": ("accept", None),
    "Yes, the whole interlude goes.": ("accept", None),
    "Author says go ahead.": ("accept", None),
    "Take it, at the number Ines had.": ("accept", None),
    "Yes, but only five lines out, not seven.": ("accept", -5),
    "Yes, but six lines, not four.": ("accept", 6),
    "Yes, and take the next three lines with it: twelve out.": ("accept", -12),
    "Cut, but six lines only.": ("accept", -6),
    "Cut, but four lines only.": ("accept", -4),
    "In, but two lines out, not four.": ("accept", -2),
    "Park until the author has read it.": ("defer", None),
    "On reflection, leave this for the author.": ("defer", None),
    "Hold it.": ("defer", None),
    "Leave it with the author.": ("defer", None),
    "No. The list of ship names stays out.": ("withdraw", None),
    "No. The author wants the interlude kept.": ("withdraw", None),
    "Withdrawn: the letter stays as drafted.": ("withdraw", None),
    "Drop it.": ("withdraw", None),
    "Not this time.": ("withdraw", None),
    "Back to Ines's call.": ("revert:Ines", None),
    "As Ines had it.": ("revert:Ines", None),
    "Ditto.": ("ditto", None),
    "Same call as ED-15.": ("same:ED-15", None),
    "Same call as ED-12.": ("same:ED-12", None),
    "Same call as ED-24.": ("same:ED-24", None),
    "Same call as ED-33.": ("same:ED-33", None),
}
STATE = {"accept": "ACCEPTED", "defer": "DEFERRED", "withdraw": "WITHDRAWN"}


def page_of(line: int) -> int:
    return (line - 1) // LINES_PER_PAGE + 1


def read_sheet(name: str, inp: Path = INP) -> list[dict]:
    """One sheet of edit_register.xlsx as header-keyed rows (dates as ISO strings)."""
    import openpyxl
    ws = openpyxl.load_workbook(inp / "edit_register.xlsx", data_only=True)[name]
    head = [c.value for c in ws[1]]
    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[0] is None:
            continue
        rows.append({h: (v.date().isoformat() if hasattr(v, "date") and callable(v.date) else v)
                     for h, v in zip(head, r)})
    return rows


def read_thread(inp: Path = INP, quotes_counted=False) -> list[dict]:
    """The decisions in edit_thread.txt, in message order, as Log-shaped rows.

    A message starts at a "From:" header; its "Date:" gives logged_on; the sender's first name is the reader;
    a decision is a body line "ED-xx: <wording>". Lines beginning with ">" are the quoted earlier message and
    are skipped (quotes_counted=True reads them as decisions: a declared wrong reading).
    """
    edits = {e["edit_code"]: e for e in read_sheet("Edits", inp)}
    rows, reader, date = [], None, None
    for raw in (inp / "edit_thread.txt").read_text(encoding="utf-8").splitlines():
        line = raw.rstrip()
        if line.startswith("From: "):
            reader, date = line[6:].split(" ")[0], None
            continue
        if line.startswith("Date: ") and date is None:
            date = line[6:].strip()
            continue
        if line.startswith(">"):
            if not quotes_counted:
                continue
            line = line.lstrip("> ")
        if line[:3] == "ED-" and ": " in line:
            code, decision = line.split(": ", 1)
            rows.append({"logged_on": date, "edit_code": code, "passage_id": edits[code]["passage_id"],
                         "proposed_change": int(edits[code]["proposed_change"]), "reader": reader, "decision": decision})
    return rows


def read_inputs(inp: Path = INP, quotes_counted=False):
    ledger = list(csv.DictReader(io.open(inp / "passage_ledger.csv", encoding="utf-8")))
    return ledger, read_thread(inp, quotes_counted)


def passages(ledger, file_order=False, stretch_separate=False, stretch_overwrite=False):
    """The passages in manuscript order as (passage_id, drafted_lines). Switches are wrong readings."""
    rows = ledger if file_order else sorted(ledger, key=lambda r: int(r["draft_line"]))
    if stretch_separate:
        return [(r["passage_id"], int(r["drafted_lines"])) for r in rows]
    out: dict[str, int] = {}
    for r in rows:
        pid, n = r["passage_id"], int(r["drafted_lines"])
        if pid in out and stretch_overwrite:
            out[pid] = n
        else:
            out[pid] = out.get(pid, 0) + n
    return list(out.items())


def standing(log, withdrawn_as_deferred=False, ignore_revisions=False, revert_noop=False,
             revert_to_proposed=False, first_line_only=False, first_read_only=False, every_line=False,
             same_copies_number=False, ditto_is_accept=False):
    """Each edit's call after the whole Log: {edit_code: {passage_id, state, change}}.

    The switches are the declared wrong readings (README section 5); all False is the gold.
    every_line returns one entry per Log line that reads as accepted (code#index), applied additively.
    """
    if first_read_only:
        second = min(l["logged_on"] for l in log if str(l["logged_on"]) >= "2026-09")
        log = [l for l in log if l["logged_on"] < second]
    calls, history, extra, line_state = {}, {}, {}, []
    for i, l in enumerate(log):
        code = l["edit_code"]
        kind, lines = DECISIONS[l["decision"]]
        if first_line_only and code in calls:
            line_state.append(calls[code]["state"])
            continue
        prev = calls.get(code)
        number = int(l["proposed_change"]) if prev is None else prev["change"]
        if kind.startswith("revert:"):
            if revert_noop:
                line_state.append(prev["state"])
                continue
            reader = kind.split(":", 1)[1]
            back = [h for h in history[code] if h["reader"] == reader][-1]
            call = dict(back["call"])
            if revert_to_proposed:
                call["change"] = int(history[code][0]["proposed"])
        else:
            if kind == "ditto":
                state = "ACCEPTED" if ditto_is_accept else line_state[-1]
                change = number
            elif kind.startswith("same:"):
                target = calls[kind.split(":", 1)[1]]
                state = target["state"]
                change = target["change"] if same_copies_number else number
            else:
                state = STATE[kind]
                change = number if (lines is None or ignore_revisions) else lines
            if state == "WITHDRAWN" and withdrawn_as_deferred:
                state = "DEFERRED"
            call = {"passage_id": l["passage_id"], "state": state, "change": int(change)}
        calls[code] = call
        line_state.append(call["state"])
        history.setdefault(code, []).append({"reader": l["reader"], "call": dict(call),
                                             "proposed": l["proposed_change"]})
        if every_line and call["state"] == "ACCEPTED":
            extra[f"{code}#{i}"] = dict(call)
    if every_line:
        for code, call in calls.items():
            if call["state"] != "ACCEPTED":
                extra[code] = call
        return extra
    return calls


def layout(order, calls, no_floor=False, floor_takes_no_line=False, deferred_counts=False,
           exclusive_end=False, geometry_first=False, row_per_edit=False):
    """Lay the passages out; returns (rows, results, notes). Switches are declared wrong readings."""
    edits: dict[str, list[dict]] = {}
    for code, c in calls.items():
        edits.setdefault(c["passage_id"], []).append(dict(c, code=code.split("#")[0]))
    units = []
    for pid, drafted in order:
        mine = edits.get(pid, [])
        if row_per_edit and len(mine) > 1:
            units += [(pid, drafted, [e]) for e in mine]
        else:
            units.append((pid, drafted, mine))
    rows, notes = [], {}
    counts = {"STRADDLES_BREAK": 0, "EDIT_DEFERRED": 0, "LENGTH_FLOORED": 0, "WHOLLY_ON_PAGE": 0}
    line = 1
    for pid, drafted, mine in units:
        counted = [e for e in mine if e["state"] == "ACCEPTED" or (deferred_counts and e["state"] == "DEFERRED")]
        raw = drafted + sum(e["change"] for e in counted)
        floored = raw < 1 and not no_floor
        length = 1 if floored else raw
        occupies = max(raw, 0) if (floored and floor_takes_no_line) else length
        first = line
        last_page = (first - 1 + length) // LINES_PER_PAGE + 1 if exclusive_end else page_of(first + length - 1)
        geo = "STRADDLES_BREAK" if page_of(first) != last_page else "WHOLLY_ON_PAGE"
        deferred = any(e["state"] == "DEFERRED" for e in mine)
        if floored:
            verdict = "LENGTH_FLOORED"
        elif geometry_first and geo == "STRADDLES_BREAK":
            verdict = geo
        elif deferred:
            verdict = "EDIT_DEFERRED"
        else:
            verdict = geo
        counts[verdict] += 1
        rows.append([pid, page_of(first), verdict])
        notes[pid] = {"first": first, "last": first + length - 1, "length": length, "drafted": drafted,
                      "edits": [f"{e['code']} {e['change']:+d} {e['state']}" for e in mine]}
        line = first + occupies
    results = {
        "straddling_passage_count": counts["STRADDLES_BREAK"],
        "deferred_edit_count": counts["EDIT_DEFERRED"],
        "length_floored_count": counts["LENGTH_FLOORED"],
        "wholly_on_page_count": counts["WHOLLY_ON_PAGE"],
        "final_page_count": page_of(line - 1),
    }
    return rows, results, notes


LEDGER_SWITCHES = {"file_order", "stretch_separate", "stretch_overwrite"}
READ_SWITCHES = {"quotes_counted"}
LOG_SWITCHES = {"withdrawn_as_deferred", "ignore_revisions", "revert_noop", "revert_to_proposed",
                "first_line_only", "first_read_only", "every_line", "same_copies_number", "ditto_is_accept"}


def solve(inp: Path = INP, **switches):
    """The whole task under one set of switches: (rows, results, notes)."""
    ledger, log = read_inputs(inp, **{k: v for k, v in switches.items() if k in READ_SWITCHES})
    order = passages(ledger, **{k: v for k, v in switches.items() if k in LEDGER_SWITCHES})
    calls = standing(log, **{k: v for k, v in switches.items() if k in LOG_SWITCHES})
    rest = {k: v for k, v in switches.items() if k not in LEDGER_SWITCHES | LOG_SWITCHES | READ_SWITCHES}
    return layout(order, calls, **rest)


def derive():
    ledger, log = read_inputs()
    by_draft = sorted(ledger, key=lambda r: int(r["draft_line"]))
    line = 1
    for r in by_draft:   # the draft is contiguous: each flag starts where the one before it ended
        assert int(r["draft_line"]) == line, f"{r['passage_id']}: draft_line {r['draft_line']}, expected {line}"
        line += int(r["drafted_lines"])
    order = passages(ledger)
    ids = [pid for pid, _ in order]
    stretches = {pid: [r for r in by_draft if r["passage_id"] == pid] for pid in ids}
    for pid, rows in stretches.items():
        idx = [by_draft.index(r) for r in rows]
        assert idx == list(range(idx[0], idx[0] + len(idx))), f"{pid}: its stretches are not adjacent in the draft"
    # the requester vouches for the Draft pages sheet ("the ids and the lines are right; the pages are the draft's")
    draft = read_sheet("Draft pages")
    want = [(r["passage_id"], int(r["draft_line"]), int(r["drafted_lines"]),
             (int(r["draft_line"]) - 1) // 30 + 1,
             "on one page" if (int(r["draft_line"]) - 1) // 30 == (int(r["draft_line"]) + int(r["drafted_lines"]) - 2) // 30
             else "over a break") for r in ledger]
    got = [(d["passage_id"], int(d["draft_line"]), int(d["drafted_lines"]), int(d["draft_page"]), d["sits"]) for d in draft]
    assert got == want, "Draft pages sheet is not the per-flag draft layout"
    dates = [l["logged_on"] for l in log]
    assert dates == sorted(dates), "the Log is not in entry order"
    missing = sorted({l["decision"] for l in log} - set(DECISIONS))
    assert not missing, f"decisions with no declared meaning: {missing}"
    history: dict[str, list[dict]] = {}
    for i, l in enumerate(log):
        assert l["passage_id"] in ids, f"{l['edit_code']}: names no ledger passage"
        kind = DECISIONS[l["decision"]][0]
        assert not (kind == "ditto" and i == 0), "Ditto. on the first line"
        if kind.startswith("same:"):
            assert kind.split(":")[1] in history, f"{l['edit_code']}: same-as an edit not yet logged"
        history.setdefault(l["edit_code"], []).append(l)
    for code, lines in history.items():
        assert len({x["passage_id"] for x in lines}) == 1, f"{code}: lines disagree on the passage"
    assert set(history) == {e["edit_code"] for e in read_sheet("Edits")}, "an edit has no decision, or a decision no edit"
    text = (INP / "edit_thread.txt").read_text(encoding="utf-8")
    assert sum(1 for l in text.splitlines() if l.startswith("> ED-")) >= 40, "the thread quotes too few decisions"
    quoted = read_thread(INP, quotes_counted=True)
    assert len(quoted) > len(log), "quoted lines add no decisions"

    calls = standing(log)
    rows, results, notes = layout(order, calls)


    # design invariants: the cases the difficulty design (README section 5) depends on
    drafted = dict(order)
    kinds = {c: [DECISIONS[x["decision"]][0] for x in lines] for c, lines in history.items()}
    assert sum(1 for s in stretches.values() if len(s) > 1) == 1, "expected exactly one two-stretch passage"
    two = next(pid for pid, s in stretches.items() if len(s) > 1)
    assert ledger.index(stretches[two][-1]) > len(ledger) - 5, f"{two}: second stretch not near the ledger's end"
    assert any(v["state"] == "WITHDRAWN" and kinds[c][0] == "accept"
               and -int(history[c][0]["proposed_change"]) == drafted[v["passage_id"]] for c, v in calls.items()), \
        "no exact-length cut accepted and later withdrawn"
    assert any(v["state"] == "DEFERRED" and kinds[c][0] == "accept"
               and -int(history[c][0]["proposed_change"]) == drafted[v["passage_id"]] for c, v in calls.items()), \
        "no exact-length cut accepted and later deferred"
    assert any(k[0] == "defer" and k[-1] == "accept" for k in kinds.values()), "no deferral later accepted"
    assert any(k[-1].startswith("revert") for k in kinds.values()), "no decision returning to an earlier call"
    assert any(k[-1].startswith("same:") and calls[c]["state"] == "WITHDRAWN" for c, k in kinds.items()), \
        "no same-as decision that withdraws"
    assert any(k[-1] == "ditto" and calls[c]["state"] != "ACCEPTED" for c, k in kinds.items()), \
        "no Ditto. that is not an acceptance"
    assert any(DECISIONS[x["decision"]][1] is not None for lines in history.values() for x in lines[1:]), \
        "no revision on a later line"
    multi = [pid for pid in drafted if sum(1 for v in calls.values()
                                           if v["passage_id"] == pid and v["state"] == "ACCEPTED") > 1]
    assert any(notes[pid]["length"] == 1 and drafted[pid] > 1 for pid in multi), \
        "no passage whose several accepted edits together floor it"
    assert any(n["last"] % LINES_PER_PAGE == 0 and n["length"] > 1 for n in notes.values()), \
        "no passage ends on a page's last line"
    return rows, results, notes


def render_csv(rows) -> str:
    out = io.StringIO()
    w = csv.writer(out, lineterminator="\n")
    w.writerow(HEADER)
    w.writerows(rows)
    return out.getvalue()


def render_json(results) -> str:
    return json.dumps(results, indent=2) + "\n"


def check_pins(rows, results) -> list[str]:
    """Every value tests/verifier.json pins must equal the derived gold."""
    spec = json.loads((ROOT / "tests" / "verifier.json").read_text(encoding="utf-8"))
    gold = {r[0]: {"page_number": str(r[1]), "verdict": r[2]} for r in rows}
    problems, graded = [], set()
    for v in spec["verifiers"]:
        exp = v["assertion"].get("expected")
        if isinstance(exp, dict) and "rows" in exp:
            for pid, cells in exp["rows"].items():
                graded.add(pid)
                if cells != gold.get(pid):
                    problems.append(f"{v['name']}: {pid} pinned {cells}, gold {gold.get(pid)}")
            if "row_set" in exp and sorted(exp["row_set"]) != sorted(gold):
                problems.append(f"{v['name']}: row_set differs from the ledger")
        if isinstance(exp, dict) and "keys" in exp:
            pinned = {k: s["value"] for k, s in exp["keys"].items()}
            if pinned != results:
                problems.append(f"{v['name']}: pins {pinned}, gold {results}")
    if graded != set(gold):
        problems.append(f"rows graded {sorted(graded)} != ledger {sorted(gold)}")
    return problems


def main() -> int:
    rows, results, notes = derive()
    csv_text, json_text = render_csv(rows), render_json(results)
    if "--write" in sys.argv:
        (GOLD / "page_register.csv").write_text(csv_text, encoding="utf-8")
        (GOLD / "results.json").write_text(json_text, encoding="utf-8")
        print("wrote solution/files/page_register.csv and results.json")
    for pid, n in notes.items():
        print(f"  {pid}: lines {n['first']}-{n['last']} ({n['length']}) edits {n['edits']}")
    print("  " + json.dumps(results))
    bad = []
    if (GOLD / "page_register.csv").read_text(encoding="utf-8") != csv_text:
        bad.append("solution/files/page_register.csv differs from the derivation")
    if (GOLD / "results.json").read_text(encoding="utf-8") != json_text:
        bad.append("solution/files/results.json differs from the derivation")
    bad += check_pins(rows, results)
    if bad:
        print("\n".join("MISMATCH " + b for b in bad))
        return 1
    print("gold agrees with the inputs, and every verifier pin agrees with the gold")
    return 0


if __name__ == "__main__":
    sys.exit(main())
